#!/usr/bin/env python3
"""Fill verified academy DOCX templates from Agent-prepared facts; never sign/approve."""
import argparse
from copy import deepcopy
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.table import _Cell
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

TEMPLATES = Path(__file__).resolve().parents[1] / 'assets' / 'templates'


def amount(value, label, positive=False):
    if not isinstance(value, str):
        raise ValueError(f'{label}: supply a decimal string, not a JSON number')
    try:
        result = Decimal(value)
    except InvalidOperation:
        raise ValueError(f'{label}: invalid decimal')
    if not result.is_finite() or result < 0 or (positive and result == 0):
        raise ValueError(f'{label}: invalid amount or quantity')
    return result


def money(value):
    return format(value.quantize(Decimal('.01'), rounding=ROUND_HALF_UP), '.2f')


def rows(table):
    return list(table._tbl.findall(qn('w:tr')))


def cell(row, index, table):
    return _Cell(row.findall(qn('w:tc'))[index], table)


def write(target, value=''):
    # Preserve the first paragraph/run formatting and all cell properties.
    p = target.paragraphs[0]
    style = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for child in list(target._tc):
        if child.tag != qn('w:tcPr') and child is not p._p:
            target._tc.remove(child)
    p.clear()
    run = p.add_run(str(value or ''))
    if style is not None:
        run._r.insert(0, style)


def expand(row):
    for height in row.findall('.//' + qn('w:trHeight')):
        height.set(qn('w:hRule'), 'atLeast')


def replace_rows(table, originals, content):
    proto = originals[0]
    parent = proto.getparent()
    index = list(parent).index(proto)
    for old in originals:
        parent.remove(old)
    for offset, values in enumerate(content):
        new = deepcopy(proto)
        expand(new)
        for i, value in enumerate(values):
            write(cell(new, i, table), value)
        parent.insert(index + offset, new)


def add_photos(target, photos, base):
    write(target)
    for i, photo in enumerate(photos):
        path = (base / photo['path']).resolve()
        if not path.is_file():
            raise ValueError(f'Missing photo: {path.name}')
        with Image.open(path) as im:
            if im.format not in ('PNG', 'JPEG'):
                raise ValueError('Convert photo to PNG/JPEG before filling')
            w, h = im.size
        scale = min(5.2 / w, 1.4 / h)
        p = target.paragraphs[0] if i == 0 else target.add_paragraph()
        p.add_run().add_picture(str(path), width=Inches(w * scale), height=Inches(h * scale))
        target.add_paragraph(photo.get('caption') or f'交付照片 {i + 1}')


def fill_goods(doc, data, base, missing):
    table = doc.tables[0]
    r = rows(table)
    items = data.get('items', [])
    if not items:
        raise ValueError('items must contain the real delivery details')
    total = Decimal('0')
    values = []
    for i, item in enumerate(items, 1):
        if not item.get('name'):
            raise ValueError(f'Item {i} needs a name')
        qty = amount(item['quantity'], 'quantity', True)
        price = amount(item['unit_price'], 'unit_price')
        subtotal = (qty * price).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
        if item.get('subtotal') is not None and amount(item['subtotal'], 'subtotal') != subtotal:
            raise ValueError(f'Item {i}: quantity × unit_price does not match subtotal; explain discounts before filling')
        total += subtotal
        values.append([str(i), item['name'], item.get('model', ''), str(qty), str(price), money(subtotal), item.get('note', '')])
    if data.get('invoice_total') is not None:
        if amount(data['invoice_total'], 'invoice_total') != total:
            raise ValueError('Invoice total differs from delivery detail total; reconcile before filling')
    else:
        missing.append('核对明细合计与真实发票金额')
    replace_rows(table, r[1:4], values)
    write(cell(r[4], 1, table), 'RMB ' + money(total))
    photos = data.get('photos', [])
    add_photos(cell(r[6], 0, table), photos, base)
    expand(r[6])
    if not photos:
        missing.append('补充真实验收照片或经学院认可的交付证据')
    if data['template'] == 'self-purchase':
        # Leave actual acceptance checks and signatures for the responsible people.
        write(cell(r[7], 1, table), '□ 合格    □ 不合格')
        if data.get('acceptance_date'):
            day = date.fromisoformat(data['acceptance_date'])
            write(cell(r[8], 1, table), f'{day.year} 年 {day.month} 月 {day.day} 日')
        else:
            missing.append('补充实际验收日期')
        missing.append('团队经办人、辅导老师核验并完成真实验收结论及签字')
    else:
        missing.append('有权验收人员填写五项检查、学院验收意见、日期及签署')
    return {'detail_total': money(total), 'currency': 'CNY'}


def fill_service(doc, data, missing):
    first, second = doc.tables
    a, b = rows(first), rows(second)
    fields = data.get('service', {})
    slots = [('project', a[0], 1, first), ('handler', a[1], 1, first),
             ('department', a[1], 3, first), ('supplier', a[2], 1, first),
             ('period', a[2], 3, first), ('project', b[0], 1, second),
             ('request_department', b[1], 1, second), ('department', b[1], 3, second),
             ('content', b[2], 1, second), ('period', b[3], 3, second),
             ('supplier', b[4], 1, second)]
    for key, row, index, table in slots:
        write(cell(row, index, table), fields.get(key, ''))
        expand(row)
        if not fields.get(key):
            missing.append('补充服务信息：' + key)
    if fields.get('contract_total') is not None:
        total = amount(fields['contract_total'], 'contract_total')
        write(cell(b[3], 1, second), money(total) + ' 元')
    else:
        write(cell(b[3], 1, second), '元')
        missing.append('补充合同价（不是本次付款额）')
    checks = fields.get('checks', [])
    values = []
    for i, check in enumerate(checks, 1):
        if not check.get('criterion'):
            raise ValueError('Each service check needs a criterion based on the real contract/delivery')
        values.append(['', f"{i}. {check['criterion']}", '通过 □    不通过 □', check.get('evidence', '')])
    if not values:
        values = [['', '', '通过 □    不通过 □', ''] for _ in range(4)]
        missing.append('按实际合同与交付补充服务验收分项')
    replace_rows(first, a[4:8], values)
    missing.append('有权人员核实服务交付并填写结论、履约评分、双方签名盖章和日期')
    return {'currency': 'CNY', 'service_checks': len(checks)}


def generate(data, output, base):
    if data.get('currency', 'CNY') != 'CNY':
        raise ValueError('Original templates use RMB; reconcile foreign currency and use a reviewed conversion before filling')
    item = next((x for x in json.loads((TEMPLATES / 'manifest.json').read_text())['templates'] if x['id'] == data['template']), None)
    if item is None:
        raise ValueError('Unknown template')
    source = (TEMPLATES / item['file']).read_bytes()
    if hashlib.sha256(source).hexdigest() != item['sha256']:
        raise ValueError('Template hash mismatch')
    output = output.resolve()
    report = output.with_suffix('.materials.json')
    if output.suffix.lower() != '.docx' or TEMPLATES in output.parents:
        raise ValueError('Use a new .docx outside the template directory')
    if output.exists() or report.exists():
        raise ValueError('Output/report exists; use a new revision filename')
    doc = Document(BytesIO(source))
    missing = list(data.get('missing', []))
    summary = fill_service(doc, data, missing) if data['template'] == 'service' else fill_goods(doc, data, base, missing)
    if data.get('render_font'):
        # Optional, explicit fallback when the source's Chinese fonts are unavailable.
        for run in doc._element.findall('.//' + qn('w:r')):
            text = ''.join(x.text or '' for x in run.findall(qn('w:t')))
            if any('\u3400' <= char <= '\u9fff' for char in text):
                properties = run.get_or_add_rPr()
                fonts = properties.find(qn('w:rFonts'))
                if fonts is None:
                    fonts = OxmlElement('w:rFonts')
                    properties.insert(0, fonts)
                for key in ('eastAsia', 'ascii', 'hAnsi'):
                    fonts.set(qn('w:' + key), data['render_font'])
    if data['template'] == 'service':
        # Supplied service template ends with an empty third section/page.
        body = doc._element.body
        tail = []
        for element in reversed(list(body)[:-1]):
            if element.tag != qn('w:p') or element.findall('.//' + qn('w:t')) or element.findall('.//' + qn('w:drawing')):
                break
            tail.append(element)
        endings = [p.find('.//' + qn('w:sectPr')) for p in reversed(tail)]
        endings = [x for x in endings if x is not None]
        if endings:
            body.replace(body[-1], deepcopy(endings[-1]))
            for element in tail:
                body.remove(element)
    # Keep the supplied package intact except the edited body and image relationships.
    generated = BytesIO()
    doc.save(generated)
    final = BytesIO()
    allowed = {'word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'}
    with ZipFile(BytesIO(source)) as original, ZipFile(generated) as changed, ZipFile(final, 'w', ZIP_DEFLATED) as result:
        for name in original.namelist():
            result.writestr(original.getinfo(name), changed.read(name) if name in allowed else original.read(name))
        for name in set(changed.namelist()) - set(original.namelist()):
            if not name.startswith('word/media/'):
                raise ValueError('Unexpected new package part: ' + name)
            result.writestr(name, changed.read(name))
    record = {'status': 'draft_needs_review_and_signatures', 'output': str(output),
              'template_id': item['id'], 'template_sha256': item['sha256'], 'summary': summary,
              'missing': list(dict.fromkeys(missing)), 'attachments': data.get('attachments', []),
              'render_font': data.get('render_font'),
              'layout_review': 'required_before_delivery'}
    os.umask(0o077)
    output.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    with output.open('xb') as stream:
        stream.write(final.getvalue())
    with report.open('x', encoding='utf-8') as stream:
        json.dump(record, stream, ensure_ascii=False, indent=2)
    return record


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    try:
        data = json.loads(args.input.read_text(encoding='utf-8'))
        result = generate(data, args.output, args.input.resolve().parent)
    except (ValueError, OSError, KeyError) as exc:
        parser.exit(1, f'{exc}\n')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
