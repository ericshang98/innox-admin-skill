"""Behavior checks for the three supplied templates; all input data is fictional."""
import importlib.util
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('fill', ROOT / 'skills/innox-invoice/scripts/fill_acceptance.py')
fill = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fill)


class AcceptanceTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.data = {'template': 'self-purchase', 'items': [
            {'name': '测试零件', 'quantity': '2', 'unit_price': '25.00'}], 'invoice_total': '50.00'}

    def generate(self, data=None, name='draft.docx'):
        path = self.base / name
        return path, fill.generate(data or self.data, path, self.base)

    def test_real_file_calculates_and_cleans_examples(self):
        path, report = self.generate()
        text = '\n'.join(c.text for t in Document(path).tables for row in t.rows for c in row.cells)
        self.assertIn('测试零件', text)
        self.assertIn('RMB 50.00', text)
        for sample in ['鼠键套装', '1950', '……']:
            self.assertNotIn(sample, text)
        self.assertIn('□ 合格    □ 不合格', text)
        self.assertEqual(len(Document(path).tables[0]._tbl.findall('.//' + fill.qn('w:sym'))), 0)
        self.assertTrue(any('照片' in x for x in report['missing']))
        self.assertTrue(any('签字' in x for x in report['missing']))
        self.assertEqual(report['status'], 'draft_needs_review_and_signatures')

    def test_mismatch_and_currency_stop_before_output(self):
        for change in [{'invoice_total': '51.00'}, {'currency': 'USD'}]:
            with self.assertRaises(ValueError):
                self.generate({**self.data, **change})
        self.assertFalse((self.base / 'draft.docx').exists())

    def test_expanded_rows_and_decimal_rounding(self):
        items = [{'name': f'测试明细 {i}', 'quantity': '3', 'unit_price': '0.335'} for i in range(12)]
        path, report = self.generate({'template': 'goods', 'items': items, 'invoice_total': '12.12'})
        self.assertEqual(report['summary']['detail_total'], '12.12')
        text = '\n'.join(c.text for row in Document(path).tables[0].rows for c in row.cells)
        for i in range(12):self.assertIn(f'测试明细 {i}', text)

    def test_photo_embedded_and_source_preserved(self):
        image = self.base / 'test-fixture.png'
        Image.new('RGB', (200, 100), 'white').save(image)
        path, _ = self.generate({**self.data, 'photos': [{'path': image.name, 'caption': 'TEST FIXTURE'}]})
        self.assertEqual(len(Document(path).inline_shapes), 1)
        original = fill.TEMPLATES / 'self-purchase.docx'
        allowed = {'word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'}
        with ZipFile(original) as a, ZipFile(path) as b:
            for part in a.namelist():
                if part not in allowed:self.assertEqual(a.read(part), b.read(part), part)
        self.assertEqual(Document(path).sections[0].page_width, Document(original).sections[0].page_width)

    def test_service_keeps_both_forms_and_no_fake_approval(self):
        path, _ = self.generate({'template': 'service', 'service': {'project': '测试项目',
            'checks': [{'criterion': '测试交付条目', 'evidence': '测试证据，不代表真实交付'}]}})
        doc = Document(path)
        self.assertEqual(len(doc.tables), 2)
        alltext = '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        for sample in ['西丽动物园研学', '梁佳荣', '行政人事', '现场搭建效率']:
            self.assertNotIn(sample, alltext)
        self.assertIn('测试交付条目', alltext)
        self.assertEqual(fill.cell(fill.rows(doc.tables[1])[10], 1, doc.tables[1]).text, '')
        self.assertEqual(fill.cell(fill.rows(doc.tables[0])[-2], 1, doc.tables[0]).text, '')

    def test_cannot_overwrite_or_generate_from_missing_photo(self):
        path, _ = self.generate()
        saved = path.read_bytes()
        with self.assertRaises(ValueError):self.generate()
        self.assertEqual(path.read_bytes(), saved)
        with self.assertRaises(ValueError):
            self.generate({**self.data, 'photos': [{'path': 'missing.png'}]}, 'other.docx')
        self.assertFalse((self.base / 'other.docx').exists())


if __name__ == '__main__':unittest.main()
